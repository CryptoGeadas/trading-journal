"""CSV and XLSX export: text stays text, coin amounts keep their decimals."""

import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

FORMULA = '=HYPERLINK("https://example.com","x")'


@pytest.fixture
def api():
    with TestClient(app) as client:
        yield client


def _add(api, **trade):
    body = {"pair": "BTC/USDT", "side": "buy", "quantity": 1, "price": 100, "timestamp": "2026-01-01T00:00:00Z", **trade}
    resp = api.post("/api/trades", json=body)
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _csv(api):
    text = api.get("/api/export", params={"format": "csv"}).text
    return list(csv.DictReader(io.StringIO(text)))


def _xlsx(api):
    data = api.get("/api/export", params={"format": "xlsx"}).content
    ws = load_workbook(io.BytesIO(data)).active
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2)]


@pytest.mark.parametrize("text", [FORMULA, "+1+1", "-2+3", "@SUM(A1)", "  =1+1"])
def test_a_note_that_looks_like_a_formula_exports_as_plain_text(api, text):
    trade_id = _add(api, notes=text, strategy="=1+1")
    api.patch(f"/api/trades/{trade_id}", json={"emotion_tag": "=cmd"})
    row = _csv(api)[0]
    assert (row["Notes"], row["Strategy"], row["Emotion"]) == ("'" + text, "'=1+1", "'=cmd")
    cells = _xlsx(api)[0]
    assert (cells["Notes"].value, cells["Notes"].data_type) == (text, "s")
    assert (cells["Strategy"].value, cells["Strategy"].data_type) == ("=1+1", "s")


def test_coin_amounts_keep_their_decimals_in_xlsx(api):
    _add(api, pair="ETH/BTC", quantity=0.5, price=0.05123456, fee=0.00004, fee_currency="BTC")
    _add(api, pair="SOL/USDT", quantity=2, price=150.125, fee=0.3, fee_currency="USDT")
    rows = {r["Pair"].value: r for r in _xlsx(api)}
    eth = rows["ETH/BTC"]
    assert eth["Total"].value == pytest.approx(0.02561728)
    assert eth["Total"].number_format == "0.00000000"
    assert (eth["Fee"].value, eth["Fee"].number_format) == (0.00004, "0.00000000")
    sol = rows["SOL/USDT"]
    assert (sol["Total"].value, sol["Total"].number_format) == (300.25, "0.00")


def test_export_has_emotion_rating_and_one_time_format(api):
    trade_id = _add(api, timestamp="2026-01-01T01:00:00+02:00")
    api.patch(f"/api/trades/{trade_id}", json={"emotion_tag": "FOMO", "trade_rating": 4})
    row = _csv(api)[0]
    assert (row["Emotion"], row["Rating"], row["Time"]) == ("FOMO", "4", "2025-12-31 23:00:00 UTC")
    cells = _xlsx(api)[0]
    assert (cells["Emotion"].value, cells["Rating"].value) == ("FOMO", 4)


# ---- #41: manual trades ----

@pytest.mark.parametrize("change, words", [
    ({"timestamp": "31/01/2026"}, "Time must look like"),
    ({"trade_type": "perp"}, "Type must be"),
    ({"fee": -1}, "Fee cannot be negative"),
])
def test_bad_manual_trade_input_is_refused_with_a_clear_message(api, change, words):
    body = {"pair": "BTC/USDT", "side": "buy", "quantity": 1, "price": 100, **change}
    resp = api.post("/api/trades", json=body)
    assert resp.status_code == 400
    assert words in resp.json()["detail"]


def test_a_manual_trade_can_be_deleted_and_a_synced_one_cannot(api):
    from app.database import get_db

    manual_id = _add(api)
    db = get_db()
    db.execute("INSERT INTO exchanges (id, exchange, api_key_enc, api_secret_enc) VALUES ('bn', 'binance', 'k', 's')")
    db.execute("""INSERT INTO trades (id, exchange_id, exchange, external_id, timestamp, pair, base_currency,
                  quote_currency, side, quantity, price, total) VALUES ('synced', 'bn', 'binance', 'x1',
                  '2026-01-02T00:00:00.000Z', 'BTC/USDT', 'BTC', 'USDT', 'sell', 1, 1, 1)""")
    db.commit()
    db.close()
    assert api.delete(f"/api/trades/{manual_id}").status_code == 400  # needs confirm
    assert api.delete(f"/api/trades/synced", params={"confirm": True}).status_code == 403
    assert api.delete(f"/api/trades/{manual_id}", params={"confirm": True}).json()["deleted"] == 1
    assert [t["id"] for t in api.get("/api/trades").json()["trades"]] == ["synced"]


def test_deleting_the_last_manual_trade_removes_its_source(api):
    trade_id = _add(api, exchange="binanec")
    assert [e["id"] for e in api.get("/api/exchanges", params={"include_manual": True}).json()] == ["manual_binanec"]
    api.delete(f"/api/trades/{trade_id}", params={"confirm": True})
    assert api.get("/api/exchanges", params={"include_manual": True}).json() == []


def test_migration_13_fixes_old_manual_futures_quotes(tmp_path):
    import sqlite3
    from app.database import MIGRATIONS

    conn = sqlite3.connect(str(tmp_path / "old.db"))
    for version in range(1, 13):
        conn.executescript(MIGRATIONS[version])
    conn.execute("INSERT INTO exchanges (id, exchange, api_key_enc, api_secret_enc) VALUES ('manual_manual', 'manual', '', '')")
    conn.execute("""INSERT INTO trades (id, exchange_id, exchange, external_id, timestamp, pair, base_currency,
                    quote_currency, side, quantity, price, total, trade_type) VALUES ('t', 'manual_manual', 'manual',
                    'x', '2026-01-01T00:00:00.000Z', 'BTC/USDT:USDT', 'BTC', 'USDT:USDT', 'buy', 1, 1, 1, 'futures')""")
    conn.commit()
    conn.executescript(MIGRATIONS[13])
    assert conn.execute("SELECT quote_currency FROM trades").fetchone() == ("USDT",)
    conn.close()


PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


def test_deleting_a_manual_order_never_touches_other_sources_with_the_same_order_id(api):
    from app.database import get_db
    a = _add(api, order_id="X1")
    _add(api, order_id="X1", timestamp="2026-01-01T00:00:01Z")
    other_manual = _add(api, exchange="otc", order_id="X1")
    db = get_db()
    db.execute("INSERT INTO exchanges (id, exchange, api_key_enc, api_secret_enc) VALUES ('bn','binance','k','s')")
    db.execute("""INSERT INTO trades (id, exchange_id, exchange, external_id, order_id, timestamp, pair, base_currency,
                  quote_currency, side, quantity, price, total) VALUES ('synced','bn','binance','e1','X1',
                  '2026-01-02T00:00:00.000Z','BTC/USDT','BTC','USDT','buy',1,1,1)""")
    db.commit()
    db.close()
    assert api.delete(f"/api/trades/{a}", params={"confirm": True}).json()["deleted"] == 2
    db = get_db()
    left = {r[0] for r in db.execute("SELECT id FROM trades").fetchall()}
    db.close()
    assert left == {"synced", other_manual}


def test_deleting_a_grouped_manual_order_removes_its_screenshots(api):
    from app import paths
    from app.database import get_db
    a = _add(api, order_id="O1")
    b = _add(api, order_id="O1", timestamp="2026-01-01T00:00:01Z")
    up = api.post(f"/api/trades/{b}/screenshots", files={"file": ("c.png", PNG, "image/png")}).json()
    stored = paths.screenshots_dir() / up["filename"]
    assert stored.exists()
    assert api.delete(f"/api/trades/{a}", params={"confirm": True}).json()["deleted"] == 2
    assert not stored.exists()
    db = get_db()
    assert db.execute("SELECT COUNT(*) FROM trade_screenshots").fetchone()[0] == 0
    db.close()


def test_a_stuck_screenshot_file_does_not_turn_a_done_delete_into_an_error(api, monkeypatch):
    from pathlib import Path
    trade_id = _add(api)
    api.post(f"/api/trades/{trade_id}/screenshots", files={"file": ("c.png", PNG, "image/png")})

    def stuck(self, missing_ok=False):
        raise PermissionError("in use")
    monkeypatch.setattr(Path, "unlink", stuck)
    resp = api.delete(f"/api/trades/{trade_id}", params={"confirm": True})
    assert resp.status_code == 200 and len(resp.json()["files_left"]) == 1


@pytest.mark.parametrize("field, raw", [("fee", '"NaN"'), ("fee", "1e400"), ("quantity", "1e400"), ("price", '"NaN"')])
def test_a_number_that_is_not_finite_is_refused(api, field, raw):
    body = '{"pair":"BTC/USDT","side":"buy","quantity":1,"price":100,"%s":%s}' % (field, raw)
    resp = api.post("/api/trades", content=body, headers={"content-type": "application/json"})
    assert resp.status_code in (400, 422)
    assert api.get("/api/trades").json()["trades"] == []


@pytest.mark.parametrize("pair", ["BTC/", "/USDT", "BTC/USDT/ETH", "BTC-USDT", "BTC/USDT:USDT"])
def test_a_malformed_or_wrong_type_pair_is_refused(api, pair):
    resp = api.post("/api/trades", json={"pair": pair, "side": "buy", "quantity": 1, "price": 1})
    assert resp.status_code == 400


def test_fees_paid_in_different_coins_are_not_added_up(api):
    from app.database import get_db
    _add(api, order_id="o1", fee=0.5, fee_currency="USDT")
    _add(api, order_id="o1", fee=0.001, fee_currency="BNB", timestamp="2026-01-01T00:00:01Z")
    row = _csv(api)[0]
    assert (row["Fee"], row["Fee Currency"]) == ("0.001 BNB; 0.5 USDT", "mixed")
    cell = _xlsx(api)[0]["Fee"]
    assert (cell.value, cell.data_type) == ("0.001 BNB; 0.5 USDT", "s")


def test_signed_numbers_stay_numbers(api):
    from app.database import get_db
    db = get_db()
    db.execute("INSERT INTO exchanges (id, exchange, api_key_enc, api_secret_enc) VALUES ('bn','binance','k','s')")
    db.execute("""INSERT INTO trades (id, exchange_id, exchange, external_id, timestamp, pair, base_currency,
                  quote_currency, side, quantity, price, total, fee, fee_currency) VALUES ('r','bn','binance','e',
                  '2026-01-02T00:00:00.000Z','BTC/USDT','BTC','USDT','buy',1,1,1,-0.01,'USDT')""")
    db.commit()
    db.close()
    assert _csv(api)[0]["Fee"] == "-0.01"
    cell = _xlsx(api)[0]["Fee"]
    assert (cell.value, cell.data_type) == (-0.01, "n")


def test_control_characters_in_notes_do_not_break_the_xlsx(api):
    _add(api, notes="pasted\x01text\x0b")
    assert _xlsx(api)[0]["Notes"].value == "pastedtext"
