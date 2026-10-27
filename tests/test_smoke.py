"""Smoke tests: every route group answers with the right shape on a demo journal.

Each test gets its own temporary data folder (conftest), loads the demo data
and calls the API like the page does. Exports are opened as real files.
"""

import csv
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


@pytest.fixture
def api(data_root):
    from app.main import app
    with TestClient(app) as client:
        assert client.post("/api/demo").status_code == 201
        yield client


def ok(resp):
    assert resp.status_code in (200, 201), (resp.request.url, resp.status_code, resp.text[:300])
    return resp.json()


def test_health_and_system(api):
    assert ok(api.get("/health"))["status"] == "ok"
    system = ok(api.get("/api/system"))
    assert {"data_dir", "left_behind"} <= set(system)


def test_overview(api):
    o = ok(api.get("/api/overview"))
    assert o["total_trades"] > 0 and o["closed_trades"] > 0 and o["currency"] == "USD"
    for key in ("total_pnl", "pnl_this_month", "pnl_this_week", "win_rate", "best_pair", "worst_pair", "total_fees"):
        assert key in o


def test_trades_group(api):
    page = ok(api.get("/api/trades", params={"page_size": 20}))
    assert page["total"] > 20 and len(page["trades"]) == 20 and page["total_pages"] > 1
    first = page["trades"][0]
    assert {"id", "pair", "side", "quantity", "price", "fill_count", "strategy", "emotion_tag"} <= set(first)
    fills = ok(api.get(f"/api/trades/fills/{first['order_id']}"))
    assert len(fills) == first["fill_count"]
    one = ok(api.get(f"/api/trades/{first['id']}"))
    assert one["id"] == first["id"]
    patched = ok(api.patch(f"/api/trades/{first['id']}", json={"notes": "checked", "trade_rating": 4}))
    assert (patched["notes"], patched["trade_rating"]) == ("checked", 4)
    assert "BTC/USDT:USDT" in ok(api.get("/api/pairs"))
    filtered = ok(api.get("/api/trades", params={"pair": "SOL/USDT", "side": "buy", "trade_type": "spot"}))
    assert filtered["total"] > 0 and all(t["pair"] == "SOL/USDT" and t["side"] == "buy" for t in filtered["trades"])
    created = ok(api.post("/api/trades", json={"pair": "ADA/USDT", "side": "buy", "quantity": 10, "price": 0.5}))
    assert ok(api.delete(f"/api/trades/{created['id']}", params={"confirm": True}))["deleted"] == 1


def test_pnl_group(api):
    for group in ("pair", "exchange", "strategy", "emotion_tag", "month", "week"):
        r = ok(api.get("/api/pnl", params={"group_by": group}))
        assert r["breakdown"] and round(sum(g["total_pnl"] for g in r["breakdown"]), 2) == r["total_pnl"]
    assert ok(api.get("/api/pnl", params={"method": "avg_cost"}))["method"] == "avg_cost"
    curve = ok(api.get("/api/equity-curve"))
    assert len(curve) > 10 and curve[-1]["cumulative_pnl"] == ok(api.get("/api/pnl"))["total_pnl"]


def test_patterns(api):
    types = {p["type"] for p in ok(api.get("/api/patterns", params={"tz": "Europe/Lisbon"}))["patterns"]}
    assert types == {"streaks", "time_of_day", "hold_time"}


def test_funding_group(api):
    fees = ok(api.get("/api/funding-fees", params={"page_size": 10}))
    assert fees["total"] > 0 and len(fees["fees"]) == 10
    summary = ok(api.get("/api/funding-fees/summary"))
    assert summary["by_pair"] and summary["total_funding"] != 0


def test_exchanges_group(api):
    assert ok(api.get("/api/exchanges")) == []  # demo accounts are not connections
    assert len(ok(api.get("/api/exchanges", params={"include_manual": True}))) == 2
    assert ok(api.get("/api/sync-log")) == []


def test_strategies_group(api):
    strategies = ok(api.get("/api/strategies"))
    assert len(strategies) == 3 and all(s["playbook"] for s in strategies)
    one = ok(api.get(f"/api/strategies/{strategies[0]['id']}"))
    assert one["name"] == strategies[0]["name"]
    exported = api.get("/api/strategies-export")
    assert exported.status_code == 200
    new = ok(api.post("/api/strategies", json={"name": "Scalp", "colour": "#34d399"}))
    ok(api.patch(f"/api/strategies/{new['id']}", json={"description": "fast"}))
    ok(api.delete(f"/api/strategies/{new['id']}", params={"confirm": True}))
    assert len(ok(api.get("/api/strategies"))) == 3


def test_emotion_tags_group(api):
    tags = ok(api.get("/api/emotion-tags"))
    assert {"FOMO", "Followed Plan"} <= {t["name"] for t in tags}
    assert api.get("/api/emotion-tags-export").status_code == 200
    new = ok(api.post("/api/emotion-tags", json={"name": "Calm"}))
    ok(api.patch(f"/api/emotion-tags/{new['id']}", json={"name": "Very calm"}))
    assert ok(api.get(f"/api/emotion-tags/{new['id']}"))["name"] == "Very calm"
    ok(api.delete(f"/api/emotion-tags/{new['id']}", params={"confirm": True}))


def test_notes_group(api):
    notes = ok(api.get("/api/notes"))
    assert len(notes) == 3
    new = ok(api.post("/api/notes", json={"title": "Today", "content": "Calm day."}))
    ok(api.patch(f"/api/notes/{new['id']}", json={"content": "Calm week."}))
    ok(api.delete(f"/api/notes/{new['id']}"))
    assert len(ok(api.get("/api/notes"))) == 3


def test_screenshots_group(api):
    trade = ok(api.get("/api/trades"))["trades"][0]
    up = ok(api.post(f"/api/trades/{trade['id']}/screenshots", files={"file": ("chart.png", PNG, "image/png")}))
    assert api.get(up["url"]).content == PNG
    assert len(ok(api.get(f"/api/trades/{trade['id']}/screenshots"))) == 1
    ok(api.delete(f"/api/screenshots/{up['id']}"))
    assert ok(api.get(f"/api/trades/{trade['id']}/screenshots")) == []


def _csv_rows(resp):
    return list(csv.DictReader(io.StringIO(resp.text)))


def test_exports_are_valid_files_that_follow_the_filters(api):
    params = {"pair": "ETH/USDT", "side": "sell", "date_from": "2026-01-01"}
    expected = ok(api.get("/api/trades", params={**params, "page_size": 200}))["total"]
    assert expected > 0

    rows = _csv_rows(api.get("/api/export", params={**params, "format": "csv"}))
    assert len(rows) == expected and {r["Pair"] for r in rows} == {"ETH/USDT"} and {r["Side"] for r in rows} == {"SELL"}

    ws = load_workbook(io.BytesIO(api.get("/api/export", params={**params, "format": "xlsx"}).content)).active
    assert ws.max_row - 1 == expected

    trade = ok(api.get("/api/trades", params={**params}))["trades"][0]
    ok(api.post(f"/api/trades/{trade['id']}/screenshots", files={"file": ("c.png", PNG, "image/png")}))
    archive = zipfile.ZipFile(io.BytesIO(api.get("/api/export", params={**params, "format": "csv",
                                                                         "include_screenshots": True}).content))
    names = archive.namelist()
    assert sum(n.endswith(".csv") for n in names) == 1 and sum(n.startswith("screenshots/") for n in names) == 1
    sheet = archive.read(next(n for n in names if n.endswith(".csv"))).decode()
    assert len(list(csv.DictReader(io.StringIO(sheet)))) == expected


def test_backup_then_restore_into_a_fresh_folder_gives_the_same_trades(api, tmp_path, monkeypatch):
    from app import encryption, paths
    before = ok(api.get("/api/trades", params={"page_size": 200}))
    pnl_before = ok(api.get("/api/pnl"))
    backup = api.get("/api/backup")
    assert backup.status_code == 200 and "journal.db" in zipfile.ZipFile(io.BytesIO(backup.content)).namelist()

    monkeypatch.setenv(paths.DATA_DIR_ENV, str(tmp_path / "fresh"))
    (tmp_path / "fresh").mkdir()
    encryption.forget_key()
    ok(api.post("/api/restore", files={"file": ("backup.zip", backup.content, "application/zip")}))
    assert ok(api.get("/api/trades", params={"page_size": 200})) == before
    assert ok(api.get("/api/pnl")) == pnl_before


def test_demo_group(api):
    assert ok(api.get("/api/demo")) == {"loaded": True}
    assert api.post("/api/demo").status_code == 409
    assert ok(api.delete("/api/demo", params={"confirm": True}))["trades"] > 0
    assert ok(api.get("/api/demo")) == {"loaded": False}


def test_the_page_itself_is_served(api):
    page = api.get("/")
    assert page.status_code == 200 and "text/html" in page.headers["content-type"]
    assert 'id="demo-banner"' in page.text


def test_strategy_and_tag_export_import_round_trip(api):
    strategies = api.get("/api/strategies-export").json()
    tags = api.get("/api/emotion-tags-export").json()
    again = ok(api.post("/api/strategies-import", json=strategies))
    assert again["imported"] == 0 and again["skipped"] == 3  # all names exist already
    ok(api.post("/api/emotion-tags-import", json=tags))
    assert len(ok(api.get("/api/strategies"))) == 3
