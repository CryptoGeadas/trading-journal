"""Export endpoints — CSV, XLSX, or ZIP (with screenshots) download."""

import io
import csv
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse
from app import paths
from app.database import get_db
from app.services.pnl_engine import currency_unit
from app.services.pnl_engine import TradeFilter

router = APIRouter()


EXPORT_COLUMNS = [
    ("timestamp", "Time"),
    ("exchange", "Exchange"),
    ("pair", "Pair"),
    ("trade_type", "Type"),
    ("side", "Side"),
    ("quantity", "Quantity"),
    ("price", "Avg Price"),
    ("total", "Total"),
    ("quote_currency", "Quote Currency"),
    ("fee", "Fee"),
    ("fee_currency", "Fee Currency"),
    ("fill_count", "Fills"),
    ("order_id", "Order ID"),
    ("strategy", "Strategy"),
    ("emotion_tag", "Emotion"),
    ("trade_rating", "Rating"),
    ("notes", "Notes"),
]

NUMBER_COLUMNS = {"quantity", "price", "total", "fee", "fill_count", "trade_rating"}
# A text cell starting with one of these is run as a formula by Excel,
# LibreOffice and Sheets (e.g. a note "=HYPERLINK(...)").
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _decimals(currency: str) -> int:
    """Dollar amounts to the cent; coin amounts (0.00004 BTC) keep 8 decimals."""
    return 2 if currency_unit(currency or "") == "USD" else 8


def _query_grouped_trades(
    exchange: Optional[str],
    pair: Optional[str],
    side: Optional[str],
    trade_type: Optional[str],
    strategy: Optional[str],
    date_from: Optional[str],
    date_to: Optional[str],
    order_by: str,
    order_dir: str,
) -> list[dict]:
    """Query trades grouped by order, applying filters."""
    db = get_db()
    try:
        where, params = TradeFilter(exchange, pair, trade_type, strategy, date_from, date_to).sql(side=side)
        group_key = "COALESCE(order_id, id)"

        query = f"""
            SELECT
                MIN(id)              as id,
                MIN(timestamp)       as timestamp,
                exchange_id,
                exchange,
                pair,
                MIN(trade_type)      as trade_type,
                side,
                SUM(quantity)        as quantity,
                CASE WHEN SUM(quantity) > 0
                     THEN SUM(total) / SUM(quantity)
                     ELSE 0 END      as price,
                SUM(total)           as total,
                SUM(fee)             as fee,
                quote_currency,
                MIN(fee_currency)    as fee_currency,
                COUNT(DISTINCT fee_currency) as fee_currency_count,
                GROUP_CONCAT(id)     as fill_ids,
                COUNT(*)             as fill_count,
                order_id,
                MIN(strategy)        as strategy,
                MIN(emotion_tag)     as emotion_tag,
                MIN(trade_rating)    as trade_rating,
                MIN(notes)           as notes
            FROM trades {where}
            GROUP BY {group_key}, exchange_id, pair, side
            ORDER BY {order_by} {order_dir}
        """
        rows = [dict(r) for r in db.execute(query, params).fetchall()]
        for row in rows:
            if row.pop("fee_currency_count") > 1:
                # Fills paid fees in different coins (e.g. USDT and BNB): never add them up.
                ids = row["fill_ids"].split(",")
                parts = db.execute(
                    f"""SELECT fee_currency, SUM(fee) FROM trades WHERE id IN ({",".join("?" * len(ids))})
                        GROUP BY fee_currency ORDER BY fee_currency""", ids).fetchall()
                row["fee"] = "; ".join(f"{amount:g} {currency or '?'}" for currency, amount in parts)
                row["fee_currency"] = "mixed"
            row.pop("fill_ids")
        return rows
    finally:
        db.close()


def _get_screenshots_for_trades(rows: list[dict]) -> dict:
    """Get screenshots for each trade/order."""
    db = get_db()
    try:
        screenshots = {}
        for row in rows:
            order_id = row.get("order_id")
            trade_id = row.get("id")

            if order_id:
                ss_rows = db.execute(
                    "SELECT * FROM trade_screenshots WHERE order_id = ? ORDER BY uploaded_at",
                    [order_id],
                ).fetchall()
                key = order_id
            else:
                ss_rows = db.execute(
                    "SELECT * FROM trade_screenshots WHERE trade_id = ? ORDER BY uploaded_at",
                    [trade_id],
                ).fetchall()
                key = trade_id

            if ss_rows:
                screenshots[key] = [dict(r) for r in ss_rows]

        return screenshots
    finally:
        db.close()


@router.get("/export")
async def export_trades(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    include_screenshots: bool = Query(False),
    exchange: Optional[str] = None,
    pair: Optional[str] = None,
    side: Optional[str] = None,
    trade_type: Optional[str] = None,
    strategy: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    order_by: str = Query("timestamp", pattern="^(timestamp|pair|exchange|total)$"),
    order_dir: str = Query("desc", pattern="^(asc|desc)$"),
):
    """Export filtered trades. Add include_screenshots=true for a ZIP with chart images."""
    rows = _query_grouped_trades(
        exchange, pair, side, trade_type, strategy, date_from, date_to, order_by, order_dir
    )

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"trades_export_{ts}"

    if include_screenshots:
        return _export_zip(rows, filename, format)
    elif format == "csv":
        return _export_csv(rows, filename)
    else:
        return _export_xlsx(rows, filename)


def _export_zip(rows: list[dict], filename: str, format: str) -> StreamingResponse:
    """Generate a ZIP containing the spreadsheet + a screenshots folder."""
    screenshots = _get_screenshots_for_trades(rows)

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        if format == "csv":
            sheet_data = _generate_csv_bytes(rows)
            zf.writestr(f"{filename}.csv", sheet_data)
        else:
            sheet_data = _generate_xlsx_bytes(rows)
            zf.writestr(f"{filename}.xlsx", sheet_data)

        ss_count = 0
        for row in rows:
            key = row.get("order_id") or row.get("id")
            if key not in screenshots:
                continue

            pair_clean = row.get("pair", "unknown").replace("/", "_")
            ts_str = row.get("timestamp", "")[:10]
            side = row.get("side", "").lower()

            for idx, ss in enumerate(screenshots[key], 1):
                src_path = paths.screenshots_dir() / ss["filename"]
                if not src_path.exists():
                    continue

                ext = Path(ss["filename"]).suffix
                archive_name = f"screenshots/{pair_clean}_{ts_str}_{side}_{idx}{ext}"
                zf.write(str(src_path), archive_name)
                ss_count += 1

        print(f"[export] ZIP created: {len(rows)} trades, {ss_count} screenshots")

    zip_buffer.seek(0)
    return StreamingResponse(
        iter([zip_buffer.getvalue()]),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename={filename}.zip"},
    )


def _generate_csv_bytes(rows: list[dict]) -> str:
    """Generate CSV content as a string."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([col[1] for col in EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([_csv_cell(row.get(key, ""), key) for key, _ in EXPORT_COLUMNS])
    return output.getvalue()


def _number(value, key: str):
    """The cell's number, or None to write the value as text (a broken row still exports)."""
    if value in (None, ""):
        return None
    try:
        return int(value) if key in ("fill_count", "trade_rating") else float(value)
    except (TypeError, ValueError):
        if not (key == "fee" and isinstance(value, str) and ";" in value):  # mixed-coin fees are text on purpose
            print(f"[export] {key} {value!r} is not a number; written as text")
        return None


def _csv_cell(value, key: str):
    if key in NUMBER_COLUMNS and (value is None or isinstance(value, (int, float))):
        return "" if value is None else value
    text = _format_value(value, key)
    # Spreadsheet apps may trim spaces first, so look past them.
    if isinstance(text, str) and text.lstrip(" \n").startswith(FORMULA_START):
        return "'" + text  # shown as text, not run
    return text


def _generate_xlsx_bytes(rows: list[dict]) -> bytes:
    """Generate XLSX content as bytes."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Trades"

    header_fill = PatternFill(start_color="1C1D24", end_color="1C1D24", fill_type="solid")
    header_text = Font(bold=True, size=11, color="E8E9ED")
    number_fmt_2 = '0.00'
    number_fmt_8 = '0.00000000'

    headers = [col[1] for col in EXPORT_COLUMNS]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_text
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (key, _) in enumerate(EXPORT_COLUMNS, 1):
            value = row.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx)
            number = _number(value, key) if key in NUMBER_COLUMNS else None
            if number is not None:
                cell.value = number
                if key == "total":
                    cell.number_format = number_fmt_2 if _decimals(row.get("quote_currency")) == 2 else number_fmt_8
                elif key == "fee":
                    cell.number_format = number_fmt_2 if _decimals(row.get("fee_currency")) == 2 else number_fmt_8
                elif key in ("quantity", "price"):
                    cell.number_format = number_fmt_8
            else:
                text = _format_value(value, key)
                if isinstance(text, str):
                    text = ILLEGAL_CHARACTERS_RE.sub("", text)  # control characters break the file
                cell.value = text
                if isinstance(text, str) and text:
                    cell.data_type = "s"  # always text, never a formula

            if key == "side" and value:
                if value.lower() == "buy":
                    cell.font = Font(color="34D399")
                elif value.lower() == "sell":
                    cell.font = Font(color="F87171")

    for col_idx in range(1, len(headers) + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row_idx in range(2, len(rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 30)

    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _export_csv(rows: list[dict], filename: str) -> StreamingResponse:
    return StreamingResponse(
        iter([_generate_csv_bytes(rows)]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
    )


def _export_xlsx(rows: list[dict], filename: str) -> StreamingResponse:
    return StreamingResponse(
        iter([_generate_xlsx_bytes(rows)]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
    )


def _format_value(value, key: str):
    """Format a value for export."""
    if value is None:
        return ""
    if key == "timestamp" and value:
        try:  # one format for synced and manual trades
            return datetime.fromisoformat(str(value).replace("Z", "+00:00")).astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        except ValueError:
            print(f"[export] time {value!r} could not be read; written as stored")
            return str(value)
    if key == "side" and value:
        return value.upper()
    if key == "exchange" and value:
        return value.capitalize()
    if key == "trade_type" and value:
        return value.capitalize()
    return value
